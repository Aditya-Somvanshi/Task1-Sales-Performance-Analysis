"""
Task 1 - Sales Performance Analysis (RaushByte Data Analytics Internship)

What this script does, step by step:
  1. Load the sales Excel file
  2. Check the data is clean
  3. Calculate total sales, top 5 products, region-wise sales, monthly sales & profit
  4. Save 4 charts (PNG) and one Excel workbook with the results + native Excel charts

Run:  python sales_analysis.py
Needs: pip install pandas openpyxl matplotlib
"""
import pandas as pd
import matplotlib
matplotlib.use("Agg")                       # draw charts to files (no window needed)
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ---------- 1. LOAD ----------
df = pd.read_excel("data/data_analytics_sales_dataset.xlsx")
df["Date"] = pd.to_datetime(df["Date"])
df["Month"] = df["Date"].dt.to_period("M").astype(str)     # e.g. "2024-01"

# ---------- 2. CHECK DATA QUALITY ----------
print("Rows:", len(df), "| Missing values:", int(df.isna().sum().sum()),
      "| Duplicate orders:", int(df["Order_ID"].duplicated().sum()))
print("Sales = Quantity x Unit_Price for every row:", bool((df.Quantity * df.Unit_Price == df.Sales).all()))
print("Profit = Sales - Expense for every row:     ", bool((df.Sales - df.Expense == df.Profit).all()))
print("Date range:", df.Date.min().date(), "to", df.Date.max().date())

# ---------- 3. ANALYSIS ----------
total_sales  = df["Sales"].sum()
total_profit = df["Profit"].sum()
margin       = total_profit / total_sales * 100

products = (df.groupby("Product")
              .agg(Units_Sold=("Quantity", "sum"), Sales=("Sales", "sum"), Profit=("Profit", "sum"))
              .sort_values("Sales", ascending=False).reset_index())
top5 = products.head(5)

regions = (df.groupby("Region")
             .agg(Orders=("Order_ID", "count"), Sales=("Sales", "sum"), Profit=("Profit", "sum"))
             .sort_values("Sales", ascending=False).reset_index())
regions["Share_%"] = (regions["Sales"] / total_sales * 100).round(2)

monthly = df.groupby("Month").agg(Sales=("Sales", "sum"), Profit=("Profit", "sum")).reset_index()
monthly["Margin_%"] = (monthly["Profit"] / monthly["Sales"] * 100).round(2)

print(f"\nTotal Sales : {total_sales:,}\nTotal Profit: {total_profit:,}\nProfit margin: {margin:.2f}%")
print("\nTop 5 products by sales:\n", top5.to_string(index=False))
print("\nRegion-wise sales:\n", regions.to_string(index=False))
print("\nBest month :", monthly.loc[monthly.Sales.idxmax(), "Month"])
print("Weakest full month:", monthly.iloc[:-1].loc[monthly.iloc[:-1].Sales.idxmin(), "Month"])

# ---------- 4. CHARTS (PNG) ----------
M = 1_000_000   # show numbers in millions so axes are easy to read
BLUE, ORANGE = "#1f4e79", "#e07b39"

plt.figure(figsize=(9, 4.5))
plt.plot(monthly["Month"], monthly["Sales"] / M, marker="o", color=BLUE, label="Sales")
plt.plot(monthly["Month"], monthly["Profit"] / M, marker="o", color=ORANGE, label="Profit")
plt.title("Monthly Sales and Profit (May 2025 is a partial month - data ends 15 May)")
plt.ylabel("Amount (millions)"); plt.xticks(rotation=45); plt.legend(); plt.grid(alpha=.3)
plt.tight_layout(); plt.savefig("charts/1_monthly_sales_profit.png", dpi=150); plt.close()

plt.figure(figsize=(7, 4))
plt.barh(top5["Product"][::-1], top5["Sales"][::-1] / M, color=BLUE)
plt.title("Top 5 Products by Sales"); plt.xlabel("Sales (millions)")
for i, v in enumerate(top5["Sales"][::-1] / M): plt.text(v + 3, i, f"{v:,.0f}M", va="center")
plt.tight_layout(); plt.savefig("charts/2_top5_products.png", dpi=150); plt.close()

plt.figure(figsize=(7, 4))
plt.bar(regions["Region"], regions["Sales"] / M, color=BLUE)
plt.title("Region-wise Sales"); plt.ylabel("Sales (millions)")
for i, v in enumerate(regions["Sales"] / M): plt.text(i, v + 3, f"{v:,.0f}M", ha="center")
plt.tight_layout(); plt.savefig("charts/3_region_sales.png", dpi=150); plt.close()

plt.figure(figsize=(5.5, 5.5))
plt.pie(regions["Sales"], labels=regions["Region"], autopct="%1.1f%%", startangle=90,
        colors=["#1f4e79", "#4f81bd", "#e07b39", "#a5a5a5"])
plt.title("Share of Total Sales by Region"); plt.tight_layout()
plt.savefig("charts/4_region_share.png", dpi=150); plt.close()

# ---------- 5. EXCEL WORKBOOK (formulas + native Excel charts) ----------
wb = Workbook()
head_font = Font(bold=True, color="FFFFFF"); head_fill = PatternFill("solid", fgColor="1F4E79")

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c); cell.font = head_font; cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")

# Sheet: Data  (raw data + a Month helper column built with a formula)
ws = wb.active; ws.title = "Data"
cols = ["Order_ID", "Date", "Customer_Name", "Region", "Product", "Quantity", "Unit_Price",
        "Sales", "Expense", "Profit", "Payment_Mode", "Customer_Rating"]
ws.append(cols + ["Month"]); style_header(ws, 1, 13)
for i, row in enumerate(df[cols].itertuples(index=False), start=2):
    ws.append(list(row[:1]) + [row[1].to_pydatetime()] + list(row[2:]) + [f'=TEXT(B{i},"yyyy-mm")'])
    ws.cell(row=i, column=2).number_format = "yyyy-mm-dd"
ws.freeze_panes = "A2"
for c in range(1, 14): ws.column_dimensions[get_column_letter(c)].width = 15
N = len(df) + 1   # last data row

# Sheet: Summary (all numbers are Excel formulas, so they can be checked)
s = wb.create_sheet("Summary")
s["A1"] = "Sales Performance Summary"; s["A1"].font = Font(bold=True, size=14)
s["A3"], s["B3"] = "Total Sales",   f"=SUM(Data!H2:H{N})"
s["A4"], s["B4"] = "Total Expense", f"=SUM(Data!I2:I{N})"
s["A5"], s["B5"] = "Total Profit",  f"=SUM(Data!J2:J{N})"
s["A6"], s["B6"] = "Profit Margin", "=B5/B3"
s["A7"], s["B7"] = "Number of Orders", f"=COUNTA(Data!A2:A{N})"
for r in (3, 4, 5): s[f"B{r}"].number_format = "#,##0"
s["B6"].number_format = "0.00%"
for r in range(3, 8): s[f"A{r}"].font = Font(bold=True)

# Top 5 products
s["A10"] = "Top 5 Products by Sales"; s["A10"].font = Font(bold=True, size=12)
s.append([]); 
for c, t in enumerate(["Rank", "Product", "Sales", "Profit"], start=1): s.cell(row=11, column=c, value=t)
style_header(s, 11, 4)
for r, p in enumerate(top5["Product"], start=12):
    s.cell(row=r, column=1, value=r - 11); s.cell(row=r, column=2, value=p)
    s.cell(row=r, column=3, value=f'=SUMIFS(Data!$H$2:$H${N},Data!$E$2:$E${N},B{r})').number_format = "#,##0"
    s.cell(row=r, column=4, value=f'=SUMIFS(Data!$J$2:$J${N},Data!$E$2:$E${N},B{r})').number_format = "#,##0"

# Region-wise
s["A19"] = "Region-wise Sales"; s["A19"].font = Font(bold=True, size=12)
for c, t in enumerate(["Region", "Sales", "Profit", "Share of Sales"], start=1): s.cell(row=20, column=c, value=t)
style_header(s, 20, 4)
for r, reg in enumerate(regions["Region"], start=21):
    s.cell(row=r, column=1, value=reg)
    s.cell(row=r, column=2, value=f'=SUMIFS(Data!$H$2:$H${N},Data!$D$2:$D${N},A{r})').number_format = "#,##0"
    s.cell(row=r, column=3, value=f'=SUMIFS(Data!$J$2:$J${N},Data!$D$2:$D${N},A{r})').number_format = "#,##0"
    s.cell(row=r, column=4, value=f"=B{r}/$B$3").number_format = "0.0%"

# Monthly
s["F10"] = "Monthly Sales and Profit"; s["F10"].font = Font(bold=True, size=12)
for c, t in enumerate(["Month", "Sales", "Profit", "Margin"], start=6): s.cell(row=11, column=c, value=t)
style_header(s, 11, 9); 
for c in range(1, 5): s.cell(row=11, column=c).fill = head_fill
for r, m in enumerate(monthly["Month"], start=12):
    s.cell(row=r, column=6, value=m)
    s.cell(row=r, column=7, value=f'=SUMIFS(Data!$H$2:$H${N},Data!$M$2:$M${N},F{r})').number_format = "#,##0"
    s.cell(row=r, column=8, value=f'=SUMIFS(Data!$J$2:$J${N},Data!$M$2:$M${N},F{r})').number_format = "#,##0"
    s.cell(row=r, column=9, value=f"=H{r}/G{r}").number_format = "0.0%"
last_m = 11 + len(monthly)
s.cell(row=last_m + 1, column=6, value="Note: 2025-05 covers only 1-15 May (partial month).").font = Font(italic=True)
for col, w in zip("ABCDEFGHI", [18, 14, 16, 16, 3, 12, 16, 16, 10]): s.column_dimensions[col].width = w

# Native Excel charts
ch = LineChart(); ch.title = "Monthly Sales and Profit"; ch.height, ch.width = 8, 18
ch.add_data(Reference(s, min_col=7, max_col=8, min_row=11, max_row=last_m), titles_from_data=True)
ch.set_categories(Reference(s, min_col=6, min_row=12, max_row=last_m)); s.add_chart(ch, "K3")

b1 = BarChart(); b1.type = "bar"; b1.title = "Top 5 Products by Sales"; b1.height, b1.width = 8, 12
b1.add_data(Reference(s, min_col=3, min_row=11, max_row=16), titles_from_data=True)
b1.set_categories(Reference(s, min_col=2, min_row=12, max_row=16)); b1.legend = None; s.add_chart(b1, "K21")

b2 = BarChart(); b2.title = "Region-wise Sales"; b2.height, b2.width = 8, 12
b2.add_data(Reference(s, min_col=2, min_row=20, max_row=24), titles_from_data=True)
b2.set_categories(Reference(s, min_col=1, min_row=21, max_row=24)); b2.legend = None; s.add_chart(b2, "K38")

wb.save("output/Task1_Sales_Analysis.xlsx")
print("\nSaved: charts/*.png and output/Task1_Sales_Analysis.xlsx")
